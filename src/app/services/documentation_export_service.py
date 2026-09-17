from __future__ import annotations

from io import BytesIO
from xml.sax.saxutils import escape as xml_escape

from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)

from docx import Document
from docx.shared import Pt


def markdown_to_docx(markdown: str) -> bytes:
    """
    Convert the project's Markdown documentation into a DOCX document.
    """

    document = Document()

    styles = document.styles

    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)

    for line in markdown.splitlines():
        stripped = line.strip()

        if not stripped:
            document.add_paragraph()
            continue

        if stripped.startswith("### "):
            document.add_heading(stripped[4:], level=3)

        elif stripped.startswith("## "):
            document.add_heading(stripped[3:], level=2)

        elif stripped.startswith("# "):
            document.add_heading(stripped[2:], level=1)

        elif stripped.startswith("- "):
            paragraph = document.add_paragraph(
                style="List Bullet"
            )
            paragraph.add_run(stripped[2:])

        else:
            document.add_paragraph(stripped)

    output = BytesIO()
    document.save(output)

    return output.getvalue()


def markdown_to_pdf(markdown: str) -> bytes:
    """
    Convert the project's Markdown documentation into a PDF document.
    """

    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=LETTER,
        rightMargin=0.65 * inch,
        leftMargin=0.65 * inch,
        topMargin=0.65 * inch,
        bottomMargin=0.65 * inch,
    )

    styles = getSampleStyleSheet()

    story = []

    for line in markdown.splitlines():
        stripped = line.strip()

        if not stripped:
            story.append(Spacer(1, 6))
            continue

        if stripped.startswith("### "):
            text = xml_escape(stripped[4:])
            story.append(
                Paragraph(
                    text,
                    styles["Heading3"],
                )
            )

        elif stripped.startswith("## "):
            text = xml_escape(stripped[3:])
            story.append(
                Paragraph(
                    text,
                    styles["Heading2"],
                )
            )

        elif stripped.startswith("# "):
            text = xml_escape(stripped[2:])
            story.append(
                Paragraph(
                    text,
                    styles["Title"],
                )
            )

        elif stripped.startswith("- "):
            text = xml_escape(stripped[2:])
            story.append(
                Paragraph(
                    f"• {text}",
                    styles["BodyText"],
                )
            )

        else:
            text = xml_escape(stripped)
            story.append(
                Paragraph(
                    text,
                    styles["BodyText"],
                )
            )

    document.build(story)

    return output.getvalue()